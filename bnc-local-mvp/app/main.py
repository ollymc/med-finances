from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import shutil
import sqlite3
import subprocess
import tempfile
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, date
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, Request, Form, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, FileResponse, RedirectResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
APP_DIR = ROOT / "app"
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "bnc.sqlite3"
DOC_DIR = DATA_DIR / "documents"
INCOMING_DIR = DATA_DIR / "incoming"
EXPORT_DIR = DATA_DIR / "exports"
CONFIG_PATH = ROOT / "config.json"
SCAN_SCRIPT = ROOT / "scripts" / "scan_wia.ps1"

for p in (DATA_DIR, DOC_DIR, INCOMING_DIR, EXPORT_DIR):
    p.mkdir(parents=True, exist_ok=True)

DEFAULT_CONFIG = {
    "app_name": "BNC Local",
    "bind_host": "127.0.0.1",
    "bind_port": 8787,
    "nas_backup_path": "",
    "backup_retention": 30,
    "scanner_name_hint": "DS-740D",
    "currency": "EUR",
}


def load_config():
    if CONFIG_PATH.exists():
        cfg = DEFAULT_CONFIG.copy()
        cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
        return cfg
    CONFIG_PATH.write_text(json.dumps(DEFAULT_CONFIG, indent=2), encoding="utf-8")
    return DEFAULT_CONFIG.copy()

CONFIG = load_config()

EXPENSE_CATEGORIES = [
    "URSSAF", "CARMF / retraite", "Redevance hôpital", "Ordre professionnel",
    "Comptabilité / conseil", "Documentation scientifique", "Abonnements / revues",
    "Repas professionnels", "Déplacements", "Péages / parking", "Communication",
    "Assurances", "Formation / congrès", "Informatique / logiciels", "Frais bancaires",
    "Fiscalité / CFE", "Fournitures", "Immobilisation", "Autres dépenses"
]
INCOME_CATEGORIES = [
    "Honoraires CPAM", "Honoraires mutuelle / AMC", "Honoraires patient",
    "Honoraires autres", "Remboursement de frais", "Autres recettes"
]

app = FastAPI(title="BNC Local")
app.mount("/static", StaticFiles(directory=APP_DIR / "static"), name="static")
templates = Jinja2Templates(directory=APP_DIR / "templates")


@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA journal_mode=WAL")
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def init_db():
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS transactions (
            id TEXT PRIMARY KEY,
            tx_date TEXT NOT NULL,
            kind TEXT NOT NULL CHECK(kind IN ('income','expense')),
            label TEXT NOT NULL,
            counterparty TEXT,
            amount REAL NOT NULL,
            category TEXT NOT NULL,
            payment_method TEXT,
            payer_type TEXT,
            encounter_token TEXT,
            professional_percent REAL DEFAULT 100,
            deductible_amount REAL,
            code_2035 TEXT,
            bank_reference TEXT,
            notes TEXT,
            source TEXT DEFAULT 'manual',
            validated INTEGER DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS documents (
            id TEXT PRIMARY KEY,
            transaction_id TEXT,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            relative_path TEXT NOT NULL,
            mime_type TEXT,
            sha256 TEXT NOT NULL,
            size_bytes INTEGER NOT NULL,
            document_type TEXT,
            document_date TEXT,
            amount REAL,
            source TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY(transaction_id) REFERENCES transactions(id) ON DELETE SET NULL
        );
        CREATE UNIQUE INDEX IF NOT EXISTS idx_documents_sha ON documents(sha256);

        CREATE TABLE IF NOT EXISTS bank_imports (
            id TEXT PRIMARY KEY,
            source_file TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            row_count INTEGER NOT NULL
        );

        CREATE TABLE IF NOT EXISTS app_events (
            seq INTEGER PRIMARY KEY AUTOINCREMENT,
            event_id TEXT UNIQUE NOT NULL,
            event_type TEXT NOT NULL,
            object_id TEXT,
            payload_json TEXT,
            created_at TEXT NOT NULL
        );
        """)


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def add_event(conn, event_type: str, object_id: Optional[str], payload: dict):
    conn.execute(
        "INSERT INTO app_events(event_id,event_type,object_id,payload_json,created_at) VALUES(?,?,?,?,?)",
        (str(uuid.uuid4()), event_type, object_id, json.dumps(payload, ensure_ascii=False), now_iso())
    )


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def save_upload(upload: UploadFile, transaction_id: Optional[str], document_type: str = "Justificatif", source: str = "upload"):
    doc_id = str(uuid.uuid4())
    suffix = Path(upload.filename or "document.bin").suffix.lower() or ".bin"
    tmp = INCOMING_DIR / f"{doc_id}{suffix}"
    with tmp.open("wb") as out:
        shutil.copyfileobj(upload.file, out)
    digest = sha256_file(tmp)
    with db() as c:
        existing = c.execute("SELECT * FROM documents WHERE sha256=?", (digest,)).fetchone()
        if existing:
            tmp.unlink(missing_ok=True)
            return existing["id"], True
        rel = Path(digest[:2]) / digest[2:4]
        dest_dir = DOC_DIR / rel
        dest_dir.mkdir(parents=True, exist_ok=True)
        stored_name = f"{digest}{suffix}"
        dest = dest_dir / stored_name
        tmp.replace(dest)
        c.execute("""INSERT INTO documents
            (id,transaction_id,original_name,stored_name,relative_path,mime_type,sha256,size_bytes,document_type,source,created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (doc_id, transaction_id, upload.filename or stored_name, stored_name, str(rel / stored_name),
             upload.content_type, digest, dest.stat().st_size, document_type, source, now_iso()))
        add_event(c, "DOCUMENT_ADDED", doc_id, {"sha256": digest, "transaction_id": transaction_id})
    return doc_id, False


def make_backup(target_root: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = target_root / f"backup-{timestamp}"
    target.mkdir(parents=True, exist_ok=False)
    (target / "documents").mkdir()

    # SQLite Online Backup API: snapshot transactionally coherent.
    snapshot = target / "bnc.sqlite3"
    src = sqlite3.connect(DB_PATH)
    dst = sqlite3.connect(snapshot)
    try:
        src.backup(dst)
    finally:
        dst.close(); src.close()

    if CONFIG_PATH.exists():
        shutil.copy2(CONFIG_PATH, target / "config.json")

    # Documents are immutable/content-addressed; copy the vault as-is.
    if DOC_DIR.exists():
        shutil.copytree(DOC_DIR, target / "documents", dirs_exist_ok=True)

    manifest = {
        "format": 1,
        "created_at": now_iso(),
        "database_sha256": sha256_file(snapshot),
        "files": []
    }
    for f in (target / "documents").rglob("*"):
        if f.is_file():
            manifest["files"].append({
                "path": str(f.relative_to(target)),
                "sha256": sha256_file(f),
                "size": f.stat().st_size,
            })
    (target / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return target


def prune_backups(root: Path):
    keep = int(CONFIG.get("backup_retention", 30))
    backups = sorted([p for p in root.glob("backup-*") if p.is_dir()], reverse=True)
    for old in backups[keep:]:
        shutil.rmtree(old, ignore_errors=True)


def dashboard_data():
    with db() as c:
        rows = c.execute("SELECT kind, COALESCE(SUM(amount),0) total FROM transactions GROUP BY kind").fetchall()
        sums = {r["kind"]: r["total"] for r in rows}
        pending = c.execute("SELECT COUNT(*) n FROM transactions WHERE validated=0").fetchone()["n"]
        missing_docs = c.execute("""SELECT COUNT(*) n FROM transactions t
            WHERE NOT EXISTS (SELECT 1 FROM documents d WHERE d.transaction_id=t.id)""").fetchone()["n"]
        recent = c.execute("SELECT * FROM transactions ORDER BY tx_date DESC, created_at DESC LIMIT 12").fetchall()
    income = sums.get("income", 0)
    expense = sums.get("expense", 0)
    return {"income": income, "expense": expense, "net": income-expense, "pending": pending, "missing_docs": missing_docs, "recent": recent}


@app.on_event("startup")
def startup():
    init_db()


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    return templates.TemplateResponse("index.html", {
        "request": request, "d": dashboard_data(), "config": CONFIG,
        "income_categories": INCOME_CATEGORIES, "expense_categories": EXPENSE_CATEGORIES
    })


@app.get("/transactions", response_class=HTMLResponse)
def transactions_page(request: Request):
    with db() as c:
        rows = c.execute("""SELECT t.*, (SELECT COUNT(*) FROM documents d WHERE d.transaction_id=t.id) doc_count
                            FROM transactions t ORDER BY tx_date DESC, created_at DESC""").fetchall()
    return templates.TemplateResponse("transactions.html", {"request": request, "rows": rows})


@app.post("/transactions/new")
def new_transaction(
    tx_date: str = Form(...), kind: str = Form(...), label: str = Form(...), amount: float = Form(...),
    category: str = Form(...), counterparty: str = Form(""), payment_method: str = Form(""),
    payer_type: str = Form(""), encounter_token: str = Form(""), professional_percent: float = Form(100),
    code_2035: str = Form(""), bank_reference: str = Form(""), notes: str = Form(""),
    attachment: Optional[UploadFile] = File(None)
):
    if kind not in ("income", "expense"):
        raise HTTPException(400, "Type invalide")
    tx_id = str(uuid.uuid4())
    deductible = amount if kind == "income" else round(amount * max(0, min(professional_percent, 100)) / 100, 2)
    with db() as c:
        c.execute("""INSERT INTO transactions
        (id,tx_date,kind,label,counterparty,amount,category,payment_method,payer_type,encounter_token,
         professional_percent,deductible_amount,code_2035,bank_reference,notes,source,created_at,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (tx_id,tx_date,kind,label,counterparty,amount,category,payment_method,payer_type,encounter_token,
         professional_percent,deductible,code_2035,bank_reference,notes,"manual",now_iso(),now_iso()))
        add_event(c, "TRANSACTION_ADDED", tx_id, {"kind": kind, "amount": amount, "category": category})
    if attachment and attachment.filename:
        save_upload(attachment, tx_id)
    return RedirectResponse("/transactions", status_code=303)


@app.post("/transactions/{tx_id}/validate")
def validate_transaction(tx_id: str):
    with db() as c:
        c.execute("UPDATE transactions SET validated=1, updated_at=? WHERE id=?", (now_iso(), tx_id))
        add_event(c, "TRANSACTION_VALIDATED", tx_id, {})
    return RedirectResponse("/transactions", status_code=303)


@app.post("/transactions/{tx_id}/attachment")
def attach_document(tx_id: str, attachment: UploadFile = File(...)):
    with db() as c:
        if not c.execute("SELECT 1 FROM transactions WHERE id=?", (tx_id,)).fetchone():
            raise HTTPException(404, "Écriture inconnue")
    save_upload(attachment, tx_id)
    return RedirectResponse("/transactions", status_code=303)


@app.get("/documents/{doc_id}")
def get_document(doc_id: str):
    with db() as c:
        row = c.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
    if not row:
        raise HTTPException(404)
    path = DOC_DIR / row["relative_path"]
    return FileResponse(path, filename=row["original_name"], media_type=row["mime_type"] or "application/octet-stream")


@app.get("/transactions/{tx_id}/documents", response_class=HTMLResponse)
def transaction_documents(request: Request, tx_id: str):
    with db() as c:
        tx = c.execute("SELECT * FROM transactions WHERE id=?", (tx_id,)).fetchone()
        docs = c.execute("SELECT * FROM documents WHERE transaction_id=? ORDER BY created_at DESC", (tx_id,)).fetchall()
    if not tx:
        raise HTTPException(404)
    return templates.TemplateResponse("documents.html", {"request": request, "tx": tx, "docs": docs})


@app.get("/bank", response_class=HTMLResponse)
def bank_page(request: Request):
    return templates.TemplateResponse("bank.html", {"request": request})


def parse_euro(v: str) -> float:
    if v is None: return 0.0
    v = str(v).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try: return float(v)
    except ValueError: return 0.0


@app.post("/bank/import")
def import_bank(file: UploadFile = File(...)):
    raw = file.file.read()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin1"):
        try:
            text = raw.decode(enc); break
        except UnicodeDecodeError: pass
    if text is None: raise HTTPException(400, "Encodage CSV non reconnu")
    sample = text[:4096]
    try: dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error: dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames: raise HTTPException(400, "CSV sans entêtes")
    aliases = {h.lower().strip(): h for h in reader.fieldnames}
    def col(*names):
        for n in names:
            if n in aliases: return aliases[n]
        return None
    c_date = col("date", "date opération", "date operation", "date de l'opération")
    c_label = col("libellé", "libelle", "description", "intitulé", "intitule")
    c_debit = col("débit", "debit")
    c_credit = col("crédit", "credit")
    c_amount = col("montant", "amount")
    if not c_date or not c_label or (not c_amount and not (c_debit or c_credit)):
        raise HTTPException(400, "Colonnes attendues: Date, Libellé, et Montant ou Débit/Crédit")
    count = 0
    with db() as c:
        for r in reader:
            label = (r.get(c_label) or "").strip()
            if not label: continue
            if c_amount:
                amt = parse_euro(r.get(c_amount))
            else:
                credit = parse_euro(r.get(c_credit)) if c_credit else 0
                debit = parse_euro(r.get(c_debit)) if c_debit else 0
                amt = credit - debit
            kind = "income" if amt >= 0 else "expense"
            category = "Honoraires autres" if kind == "income" else "Autres dépenses"
            tx_id = str(uuid.uuid4())
            raw_date = (r.get(c_date) or "").strip()
            # Keep raw date if parsing fails; normalize common FR dates.
            tx_date = raw_date
            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                try: tx_date = datetime.strptime(raw_date, fmt).date().isoformat(); break
                except ValueError: pass
            c.execute("""INSERT INTO transactions
                (id,tx_date,kind,label,amount,category,professional_percent,deductible_amount,source,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (tx_id,tx_date,kind,label,abs(amt),category,100,abs(amt),"bank_csv",now_iso(),now_iso()))
            add_event(c, "BANK_TRANSACTION_IMPORTED", tx_id, {"source_file": file.filename})
            count += 1
        c.execute("INSERT INTO bank_imports(id,source_file,imported_at,row_count) VALUES(?,?,?,?)",
                  (str(uuid.uuid4()), file.filename or "bank.csv", now_iso(), count))
    return RedirectResponse("/transactions", status_code=303)


@app.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request):
    return templates.TemplateResponse("settings.html", {"request": request, "config": CONFIG})


@app.post("/settings")
def save_settings(nas_backup_path: str = Form(""), backup_retention: int = Form(30)):
    global CONFIG
    CONFIG["nas_backup_path"] = nas_backup_path.strip()
    CONFIG["backup_retention"] = max(1, min(365, backup_retention))
    CONFIG_PATH.write_text(json.dumps(CONFIG, indent=2, ensure_ascii=False), encoding="utf-8")
    return RedirectResponse("/settings", status_code=303)


@app.post("/backup")
def backup_now():
    root = CONFIG.get("nas_backup_path", "").strip()
    if not root:
        raise HTTPException(400, "Configurez d'abord le chemin NAS dans Paramètres")
    target_root = Path(root)
    try:
        target_root.mkdir(parents=True, exist_ok=True)
        target = make_backup(target_root)
        prune_backups(target_root)
        return RedirectResponse(f"/settings?backup={target.name}", status_code=303)
    except Exception as e:
        raise HTTPException(500, f"Échec sauvegarde NAS: {e}")


@app.post("/scan")
def scan_document():
    if os.name != "nt":
        raise HTTPException(501, "Le scan WIA/TWAIN est disponible sous Windows uniquement")
    out = INCOMING_DIR / f"scan-{uuid.uuid4()}.jpg"
    cmd = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(SCAN_SCRIPT), "-OutputPath", str(out)]
    try:
        subprocess.run(cmd, check=True, timeout=180)
    except subprocess.CalledProcessError as e:
        raise HTTPException(500, f"Échec du scan: {e}")
    if not out.exists():
        raise HTTPException(500, "Le scanner n'a produit aucun fichier")
    digest = sha256_file(out)
    rel = Path(digest[:2]) / digest[2:4]
    dest_dir = DOC_DIR / rel; dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{digest}.jpg"
    if not dest.exists(): out.replace(dest)
    else: out.unlink(missing_ok=True)
    doc_id = str(uuid.uuid4())
    with db() as c:
        existing = c.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()
        if existing: return RedirectResponse("/", status_code=303)
        c.execute("""INSERT INTO documents
            (id,original_name,stored_name,relative_path,mime_type,sha256,size_bytes,document_type,source,created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (doc_id,dest.name,dest.name,str(rel/dest.name),"image/jpeg",digest,dest.stat().st_size,"Scan reçu/facture","WIA/TWAIN",now_iso()))
        add_event(c,"DOCUMENT_SCANNED",doc_id,{"sha256":digest})
    return RedirectResponse("/", status_code=303)


def all_transactions():
    with db() as c:
        return c.execute("SELECT * FROM transactions ORDER BY tx_date, created_at").fetchall()


@app.get("/export/excel")
def export_excel():
    rows = all_transactions()
    wb = Workbook(); ws = wb.active; ws.title = "Journal"
    headers = ["Date","Type","Libellé","Tiers","Montant","Catégorie","Paiement","Payer","Token","Pro %","Déductible","2035","Réf banque","Validé"]
    ws.append(headers)
    for r in rows:
        ws.append([r["tx_date"],r["kind"],r["label"],r["counterparty"],r["amount"],r["category"],r["payment_method"],r["payer_type"],r["encounter_token"],r["professional_percent"],r["deductible_amount"],r["code_2035"],r["bank_reference"],r["validated"]])
    path = EXPORT_DIR / f"journal-{date.today().isoformat()}.xlsx"; wb.save(path)
    return FileResponse(path, filename=path.name)


@app.get("/export/pdf")
def export_pdf():
    rows = all_transactions(); path = EXPORT_DIR / f"journal-{date.today().isoformat()}.pdf"
    c = canvas.Canvas(str(path), pagesize=A4); w,h=A4; y=h-40
    c.setFont("Helvetica-Bold", 14); c.drawString(40,y,"BNC Local - Journal comptable"); y-=25
    c.setFont("Helvetica", 8)
    for r in rows:
        line=f"{r['tx_date']} | {r['kind']} | {r['label'][:35]} | {r['amount']:.2f} EUR | {r['category'][:25]}"
        c.drawString(40,y,line); y-=12
        if y<40: c.showPage(); c.setFont("Helvetica",8); y=h-40
    c.save(); return FileResponse(path, filename=path.name)


@app.get("/export/word")
def export_word():
    rows = all_transactions(); path = EXPORT_DIR / f"journal-{date.today().isoformat()}.docx"
    doc=Document(); doc.add_heading("BNC Local - Journal comptable", 0)
    table=doc.add_table(rows=1, cols=6); hdr=table.rows[0].cells
    for i,v in enumerate(["Date","Type","Libellé","Montant","Catégorie","2035"]): hdr[i].text=v
    for r in rows:
        cells=table.add_row().cells
        vals=[r["tx_date"],r["kind"],r["label"],f"{r['amount']:.2f} €",r["category"],r["code_2035"] or ""]
        for i,v in enumerate(vals): cells[i].text=str(v)
    doc.save(path); return FileResponse(path, filename=path.name)


@app.get("/health")
def health():
    return JSONResponse({"status":"ok","database":str(DB_PATH),"time":now_iso()})

@app.get("/documents", response_class=HTMLResponse)
def documents_inbox(request: Request):
    with db() as c:
        docs = c.execute("SELECT * FROM documents ORDER BY created_at DESC").fetchall()
        txs = c.execute("SELECT id,tx_date,label,amount FROM transactions ORDER BY tx_date DESC, created_at DESC LIMIT 500").fetchall()
    return templates.TemplateResponse("documents_inbox.html", {"request": request, "docs": docs, "txs": txs})


@app.post("/documents/{doc_id}/link")
def link_document(doc_id: str, transaction_id: str = Form(...)):
    with db() as c:
        d = c.execute("SELECT 1 FROM documents WHERE id=?", (doc_id,)).fetchone()
        t = c.execute("SELECT 1 FROM transactions WHERE id=?", (transaction_id,)).fetchone()
        if not d or not t:
            raise HTTPException(404, "Document ou écriture introuvable")
        c.execute("UPDATE documents SET transaction_id=? WHERE id=?", (transaction_id, doc_id))
        add_event(c, "DOCUMENT_LINKED", doc_id, {"transaction_id": transaction_id})
    return RedirectResponse("/documents", status_code=303)
